import streamlit as st
import pandas as pd
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import io
import datetime

# --- 設定頁面配置 ---
st.set_page_config(page_title="Meta 廣告全策略配置工具", layout="wide")

# --- 預設資料 ---
DEFAULT_TAGS = [
    "Facebook Page Admins (粉專管理員)", "Small business (小型企業)", "Entrepreneurship (創業)",
    "商業 / 生產力軟體和應用程式", "Digital marketing (數位行銷)", "Management (管理)",
    "企業顧問 (Business consultant)", "專業商業技能 (Professional business skills)",
    "Web design (網頁設計)", "Business travelers (商務旅客)", "Coworking", "Software"
]

# --- 核心函式：設定 Word 中文字體 (微軟正黑體) 與顏色 ---
# 修正點：新增 color 參數，預設為 None
def set_font_style(run, font_name='Microsoft JhengHei', size=10, bold=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    
    # 如果有傳入顏色，則設定顏色
    if color:
        run.font.color.rgb = color
        
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), font_name)

# --- 核心函式：生成 Word 表格報告 ---
def generate_docx_report(campaigns):
    doc = Document()
    
    # 文件標題
    head = doc.add_heading(level=0)
    run = head.add_run(f"Meta 廣告投放配置指令單")
    set_font_style(run, size=18, bold=True)
    
    doc.add_paragraph(f"生成日期: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}")
    doc.add_paragraph("---")

    if not campaigns:
        doc.add_paragraph("目前沒有設定任何行銷活動。")
        return doc

    for c_idx, campaign in enumerate(campaigns):
        # 1. Campaign 標題
        h1 = doc.add_heading(level=1)
        run = h1.add_run(f"行銷活動 #{c_idx+1}: {campaign['name']}")
        set_font_style(run, size=14, bold=True)

        # 2. Campaign 摘要表格 (Key-Value 形式)
        table_camp = doc.add_table(rows=2, cols=4)
        table_camp.style = 'Table Grid'
        
        # 表頭
        cells = table_camp.rows[0].cells
        headers = ["行銷活動目標", "預算類型", "預算金額", "CBO 最佳化"]
        for i, text in enumerate(headers):
            run = cells[i].paragraphs[0].add_run(text)
            set_font_style(run, bold=True)
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            # 設定背景色 (灰色)
            tcPr = cells[i]._element.tcPr
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:fill'), 'D9D9D9') # 淺灰
            tcPr.append(shd)

        # 內容
        vals = table_camp.rows[1].cells
        c_data = [
            campaign['objective'],
            campaign['budget_type'],
            f"NT$ {campaign['budget_amount']}",
            "開啟 ✅" if campaign['is_cbo'] else "關閉 ❌"
        ]
        for i, text in enumerate(c_data):
            run = vals[i].paragraphs[0].add_run(str(text))
            set_font_style(run)
            vals[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph("") # 空行

        # 3. Ad Sets 清單表格
        if campaign['ad_sets']:
            h2 = doc.add_heading(level=2)
            run = h2.add_run("廣告組合 (受眾) 詳細配置")
            set_font_style(run, size=12, bold=True)

            # 建立表格：名稱 | 轉換/版位 | 受眾詳情 | 素材 ID
            table_ads = doc.add_table(rows=1, cols=4)
            table_ads.style = 'Table Grid'
            table_ads.autofit = False 
            
            # 設定欄寬 (依比例)
            table_ads.columns[0].width = Inches(1.2) # 名稱
            table_ads.columns[1].width = Inches(1.5) # 設定
            table_ads.columns[2].width = Inches(2.5) # 受眾
            table_ads.columns[3].width = Inches(1.5) # 素材

            # 表頭設定
            hdr_cells = table_ads.rows[0].cells
            ad_headers = ["組合名稱", "基礎設定", "受眾鎖定 (Targeting)", "素材 ID (Ads)"]
            for i, text in enumerate(ad_headers):
                run = hdr_cells[i].paragraphs[0].add_run(text)
                set_font_style(run, bold=True)
                hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                # 設定背景色
                tcPr = hdr_cells[i]._element.tcPr
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:fill'), 'E6E6E6')
                tcPr.append(shd)

            # 填入 Ad Sets 資料
            for ad_set in campaign['ad_sets']:
                row_cells = table_ads.add_row().cells
                
                # Col 1: 名稱
                run = row_cells[0].paragraphs[0].add_run(ad_set['name'] if ad_set['name'] else "(未命名)")
                set_font_style(run, bold=True)
                
                # Col 2: 基礎設定
                p2 = row_cells[1].paragraphs[0]
                run = p2.add_run(f"• 目標: {ad_set['goal']}\n")
                set_font_style(run)
                run = p2.add_run(f"• 版位: {ad_set['placements_type']}\n")
                set_font_style(run)
                if ad_set['excluded_placements']:
                    run = p2.add_run(f"  (指定: {','.join(ad_set['excluded_placements'])})")
                    set_font_style(run, size=9)
                
                # Col 3: 受眾
                p3 = row_cells[2].paragraphs[0]
                # 人口統計
                run = p3.add_run(f"【人口】{ad_set['age_min']}-{ad_set['age_max']}歲 / {ad_set['gender']} / {','.join(ad_set['locations'])}\n")
                set_font_style(run)
                # 興趣
                tags = ad_set['tags'] + ad_set['manual_tags']
                tags_str = ", ".join(tags) if tags else "無"
                run = p3.add_run(f"【興趣】{tags_str}\n")
                set_font_style(run)
                # 自訂受眾 (顯示為藍色)
                if ad_set['custom_audience']:
                    run = p3.add_run(f"【自訂】{ad_set['custom_audience']}\n")
                    set_font_style(run, color=RGBColor(0, 50, 150))
                # Advantage+
                run = p3.add_run(f"【Advantage+】{'開啟' if ad_set['advantage_audience'] else '關閉'}")
                set_font_style(run)

                # Col 4: 素材
                p4 = row_cells[3].paragraphs[0]
                final_ads = ad_set['selected_ads'] + ad_set['manual_ads']
                if final_ads:
                    for ad_id in final_ads:
                        run = p4.add_run(f"□ {ad_id}\n")
                        set_font_style(run)
                else:
                    # 未指定 (顯示為紅色)
                    run = p4.add_run("(未指定)")
                    set_font_style(run, color=RGBColor(200, 0, 0))

        doc.add_paragraph("\n") # 每個 Campaign 間隔

    return doc

# --- 核心函式：解析 Word 檔 (讀取 ID 用) ---
def parse_ad_ids_from_docx(uploaded_file):
    ad_ids = []
    try:
        doc = Document(uploaded_file)
        all_paragraphs = []
        all_paragraphs.extend(doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    all_paragraphs.extend(cell.paragraphs)
        
        for i, para in enumerate(all_paragraphs):
            text = para.text.strip()
            if "【廣告組合 ID】" in text:
                parts = text.split("【廣告組合 ID】")
                candidate = ""
                if len(parts) > 1:
                    candidate = parts[1].strip().lstrip(",").lstrip(":").lstrip("，").lstrip("：").strip()
                if candidate:
                    ad_ids.append(candidate)
                else:
                    if i + 1 < len(all_paragraphs):
                        next_text = all_paragraphs[i + 1].text.strip()
                        if next_text:
                            clean_next = next_text.lstrip(",").lstrip(":").lstrip("，").lstrip("：").strip()
                            if clean_next:
                                ad_ids.append(clean_next)
    except Exception as e:
        st.error(f"解析 Word 檔時發生錯誤: {e}")
    return sorted(list(set(ad_ids)))

# --- 輔助函式：解析 CSV/Excel ---
def parse_tags_from_csv(uploaded_file):
    tags = []
    try:
        if uploaded_file.name.endswith('.csv'):
            df = pd.read_csv(uploaded_file)
        else:
            df = pd.read_excel(uploaded_file)
        if "受眾標籤 (Tag)" in df.columns:
            tags = df["受眾標籤 (Tag)"].dropna().unique().tolist()
    except Exception as e:
        st.error(f"解析受眾標籤檔時發生錯誤: {e}")
    return tags

# --- 初始化 Session State ---
if 'campaigns' not in st.session_state:
    st.session_state.campaigns = []

# --- 側邊欄 ---
with st.sidebar:
    st.header("📂 資料來源庫")
    docx_file = st.file_uploader("上傳上刊文件 (Word)", type=['docx'])
    available_ad_ids = []
    if docx_file:
        available_ad_ids = parse_ad_ids_from_docx(docx_file)
        if available_ad_ids:
            st.success(f"已抓取 {len(available_ad_ids)} 個廣告 ID")
            with st.expander("查看 ID 列表"):
                st.write(available_ad_ids)
    
    tag_file = st.file_uploader("上傳受眾標籤 (CSV/Excel)", type=['csv', 'xlsx'])
    available_tags = DEFAULT_TAGS.copy()
    if tag_file:
        uploaded_tags = parse_tags_from_csv(tag_file)
        if uploaded_tags:
            available_tags = uploaded_tags
            st.success(f"已載入 {len(available_tags)} 個標籤")

# --- 主畫面 ---
st.title("🚀 Meta 廣告全策略配置工具")
st.info("💡 支援 Multi-Campaign 架構。匯出時將生成排版清晰的 Word 表格。")

# --- 全局控制 ---
if st.button("➕ 建立一個新的行銷活動 (Campaign)", type="primary"):
    st.session_state.campaigns.append({
        "name": "", "objective": "銷售 (Sales)", "is_cbo": True,
        "budget_type": "單日預算", "budget_amount": 1000, "start_date": datetime.date.today(),
        "ad_sets": []
    })

st.markdown("---")

# --- 迴圈渲染每個 Campaign ---
campaigns_to_remove = []

for c_idx, campaign in enumerate(st.session_state.campaigns):
    with st.container():
        st.markdown(f"## 📢 行銷活動 #{c_idx + 1}")
        
        # Campaign 設定
        with st.expander(f"設定活動內容：{campaign['name'] if campaign['name'] else '(未命名)'}", expanded=True):
            col_del_camp, _ = st.columns([1, 6])
            if col_del_camp.button(f"🗑️ 刪除整個活動 #{c_idx + 1}", key=f"del_cmp_{c_idx}"):
                campaigns_to_remove.append(c_idx)

            col1, col2 = st.columns(2)
            with col1:
                campaign['name'] = st.text_input("行銷活動名稱", value=campaign['name'], placeholder="例如：2025_TOF_流量開發", key=f"cmp_name_{c_idx}")
                campaign['objective'] = st.selectbox("行銷活動目標", ["銷售 (Sales)", "潛在客戶 (Leads)", "流量 (Traffic)", "互動 (Engagement)", "知名度 (Awareness)", "應用程式推廣 (App Promotion)"], key=f"cmp_obj_{c_idx}")
                campaign['is_cbo'] = st.toggle("開啟 CBO 預算最佳化", value=campaign['is_cbo'], key=f"cmp_cbo_{c_idx}")
            with col2:
                campaign['budget_type'] = st.selectbox("預算類型", ["單日預算", "總預算"], key=f"cmp_btype_{c_idx}")
                campaign['budget_amount'] = st.number_input("預算金額 (TWD)", min_value=0, value=campaign['budget_amount'], step=100, key=f"cmp_bamt_{c_idx}")
                campaign['start_date'] = st.date_input("開始日期", value=campaign['start_date'], key=f"cmp_date_{c_idx}")

        # --- Ad Sets 管理 ---
        st.markdown(f"**👇 設定活動 #{c_idx + 1} 的廣告組合 (受眾)**")
        if st.button(f"➕ 新增廣告組合", key=f"add_adset_{c_idx}"):
            campaign['ad_sets'].append({
                "name": "", "goal": "網站", "tags": [], "manual_tags": [],
                "custom_audience": "", "age_min": 25, "age_max": 55,
                "locations": ["台灣"], "gender": "所有性別",
                "advantage_audience": True, "placements_type": "Advantage+ (自動版位)",
                "excluded_placements": [], "placement_note": "", "selected_ads": [], "manual_ads": []
            })

        adsets_to_remove = []
        for a_idx, ad_set in enumerate(campaign['ad_sets']):
            st.info(f"廣告組合 {c_idx+1}-{a_idx+1}: {ad_set['name'] if ad_set['name'] else '設定中...'}")
            c1, c2 = st.columns(2)
            with c1:
                ad_set['name'] = st.text_input("組合名稱", value=ad_set['name'], key=f"c{c_idx}_a{a_idx}_name")
                ad_set['goal'] = st.selectbox("轉換位置", ["網站", "應用程式", "訊息", "通話"], key=f"c{c_idx}_a{a_idx}_goal")
                ad_set['tags'] = st.multiselect("興趣標籤", options=available_tags, default=list(set(ad_set['tags']) & set(available_tags)), key=f"c{c_idx}_a{a_idx}_tags")
                extra = st.text_input("手動輸入標籤", key=f"c{c_idx}_a{a_idx}_extra")
                ad_set['manual_tags'] = [t.strip() for t in extra.split(",") if t.strip()]
                ad_set['advantage_audience'] = st.checkbox("開啟 Advantage+ 受眾", value=ad_set['advantage_audience'], key=f"c{c_idx}_a{a_idx}_aa")
            with c2:
                if st.button("🗑️", key=f"del_adset_c{c_idx}_a{a_idx}"):
                    adsets_to_remove.append(a_idx)
                ad_set['custom_audience'] = st.text_area("自訂受眾", value=ad_set['custom_audience'], height=68, key=f"c{c_idx}_a{a_idx}_ca")
                col_age1, col_age2 = st.columns(2)
                ad_set['age_min'] = col_age1.number_input("最小年齡", 13, 65, ad_set['age_min'], key=f"c{c_idx}_a{a_idx}_amin")
                ad_set['age_max'] = col_age2.number_input("最大年齡", 13, 65, ad_set['age_max'], key=f"c{c_idx}_a{a_idx}_amax")
                ad_set['gender'] = st.radio("性別", ["所有性別", "男性", "女性"], horizontal=True, key=f"c{c_idx}_a{a_idx}_gen")
                
                if available_ad_ids:
                    ad_set['selected_ads'] = st.multiselect("選擇素材 ID", options=available_ad_ids, default=list(set(ad_set['selected_ads']) & set(available_ad_ids)), key=f"c{c_idx}_a{a_idx}_ads")
                else:
                    man_ads = st.text_area("手動輸入素材 ID", key=f"c{c_idx}_a{a_idx}_mads")
                    ad_set['manual_ads'] = [x.strip() for x in man_ads.split("\n") if x.strip()]

            with st.expander("版位設定 (進階)", expanded=False):
                ad_set['placements_type'] = st.radio("版位模式", ["Advantage+ (自動)", "手動指定"], key=f"c{c_idx}_a{a_idx}_pmode")
                if ad_set['placements_type'] == "手動指定":
                    ad_set['excluded_placements'] = st.multiselect("排除/指定版位", ["FB動態", "IG動態", "Stories", "Reels"], key=f"c{c_idx}_a{a_idx}_excl")
        
        for i in sorted(adsets_to_remove, reverse=True):
            del campaign['ad_sets'][i]
            
    st.markdown("---")

for i in sorted(campaigns_to_remove, reverse=True):
    del st.session_state.campaigns[i]

# --- 輸出報告 (Word) ---
st.header("📄 產出投放 Brief")
st.markdown("完成所有設定後，點擊下方按鈕下載 Word 格式的指令單。")

if st.button("下載 Word 投放指令單 (.docx)", type="primary"):
    # 產生 Word 物件
    doc = generate_docx_report(st.session_state.campaigns)
    
    # 轉為二進位串流
    bio = io.BytesIO()
    doc.save(bio)
    
    st.download_button(
        label="📥 點擊下載檔案",
        data=bio.getvalue(),
        file_name=f"Meta_Brief_{datetime.date.today()}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
